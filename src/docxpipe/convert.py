"""DOCX → Markdown: структура документа и отбор рисунков/схем."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
VML_NS = "urn:schemas-microsoft-com:vml"
R_EMBED = qn("r:embed")
R_ID = qn("r:id")

_NULL_REL_RE = re.compile(
    rb"<Relationship\b[^>]*Target=(?:\"NULL\"|'NULL')[^>]*/>",
    re.I,
)

LABEL_RE = re.compile(
    r"^(?:рисунок|рис\.?|схема|диаграмма|график|чертёж|чертеж)\b",
    re.IGNORECASE,
)
REF_RE = re.compile(
    r"(?:"
    r"на\s+рисунк\w*"
    r"|на\s+рис\."
    r"|как\s+показан\w*\s+на\s+схем\w*"
    r"|на\s+схем\w*"
    r"|показан\w*\s+на\s+(?:рисунк\w*|схем\w*)"
    r"|см\.\s*(?:рис\.?|рисунк\w*|схем\w*)"
    r"|привед\w+\s+на\s+(?:рисунк\w*|схем\w*)"
    r")",
    re.IGNORECASE,
)

CONTENT_EXT = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/jpg": ".jpg",
    "image/gif": ".gif",
    "image/bmp": ".bmp",
    "image/tiff": ".tif",
    "image/x-emf": ".emf",
    "image/emf": ".emf",
    "image/x-wmf": ".wmf",
    "image/wmf": ".wmf",
    "image/svg+xml": ".svg",
    "image/webp": ".webp",
}
NO_PREVIEW = {".emf", ".wmf", ".svg"}
_INVALID_NAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


@dataclass
class ImageRecord:
    id: str
    file: str
    path: str
    sha256: str
    nbytes: int
    heading: str
    caption: str
    reason: str
    preview: bool

    def as_dict(self, *, ocr_text: str | None = None) -> dict[str, Any]:
        data = {
            "id": self.id,
            "file": self.file,
            "path": self.path,
            "sha256": self.sha256,
            "bytes": self.nbytes,
            "heading": self.heading,
            "caption": self.caption,
            "reason": self.reason,
            "preview": self.preview,
        }
        if ocr_text is not None:
            data["ocr_text"] = ocr_text
        return data


@dataclass
class ConvertResult:
    markdown: str
    media_dirname: str
    images: list[ImageRecord] = field(default_factory=list)
    blobs: dict[str, bytes] = field(default_factory=dict)
    saved: int = 0
    skipped: int = 0
    headings: int = 0
    tables: int = 0


def media_dirname(stem: str) -> str:
    """Имя папки картинок, привязанное к имени исходного файла."""
    cleaned = _INVALID_NAME.sub("_", stem).strip(" .")
    cleaned = cleaned or "document"
    return f"image_{cleaned}"


def is_figure_label(text: str) -> bool:
    return bool(LABEL_RE.match(_norm(text)))


def is_text_reference(text: str) -> bool:
    return bool(REF_RE.search(_norm(text)))


def _norm(text: str) -> str:
    s = (text or "").replace("\xa0", " ").replace("\u00ad", "")
    return re.sub(r"[ \t\r\n]+", " ", s).strip()


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def load_document(source: Path) -> Document:
    try:
        return Document(str(source))
    except KeyError:
        return Document(_sanitize_package(source))


def _sanitize_package(source: Path) -> io.BytesIO:
    buf = io.BytesIO()
    with zipfile.ZipFile(source, "r") as src, zipfile.ZipFile(
        buf, "w", compression=zipfile.ZIP_DEFLATED
    ) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename.endswith(".rels") and b"NULL" in data:
                data = _NULL_REL_RE.sub(b"", data)
            dst.writestr(info, data)
    buf.seek(0)
    return buf


def heading_level(paragraph: Paragraph) -> int | None:
    style = paragraph.style
    name = (style.name or "").strip() if style is not None else ""
    if name.lower().startswith("toc"):
        return None
    match = re.match(r"(?:Heading|Заголовок|Уровень)\s*(\d+)$", name, re.I)
    if match:
        return max(1, min(6, int(match.group(1))))
    outline = _outline_level(paragraph)
    if outline is not None and 0 <= outline <= 5:
        return outline + 1
    return None


def _outline_level(paragraph: Paragraph) -> int | None:
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.outlineLvl is not None:
        try:
            return int(p_pr.outlineLvl.val)
        except (TypeError, ValueError):
            return None
    return None


def infer_heading_level(text: str) -> int | None:
    stripped = text.strip()
    if re.match(r"^Приложение\s+", stripped, re.I):
        return 1
    if re.match(r"^Форма\s+\d+", stripped, re.I):
        return 2
    return None


def paragraph_text(paragraph: Paragraph) -> str:
    chunks: list[str] = []
    for node in paragraph._p.findall(".//" + qn("w:t")):
        parent = node.getparent()
        deleted = False
        while parent is not None:
            if parent.tag == qn("w:del"):
                deleted = True
                break
            parent = parent.getparent()
        if not deleted and node.text:
            chunks.append(node.text)
    return _norm("".join(chunks))


def _is_list(paragraph: Paragraph) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is not None and p_pr.numPr is not None:
        return True
    style = paragraph.style
    name = (style.name or "") if style is not None else ""
    return name.lower().startswith("list")


def _blip_ids(element) -> list[str]:
    found: list[str] = []
    for blip in element.findall(f".//{{{A_NS}}}blip"):
        rid = blip.get(R_EMBED)
        if rid and rid not in found:
            found.append(rid)
    for image in element.findall(f".//{{{VML_NS}}}imagedata"):
        rid = image.get(R_ID)
        if rid and rid not in found:
            found.append(rid)
    return found


def _load_images(doc: Document) -> dict[str, tuple[str, str, bytes]]:
    images: dict[str, tuple[str, str, bytes]] = {}
    for rel in doc.part.rels.values():
        if "image" not in rel.reltype:
            continue
        part = rel.target_part
        images[rel.rId] = (Path(rel.target_ref).name, part.content_type, part.blob)
    return images


def _extension(name: str, content_type: str) -> str:
    ext = CONTENT_EXT.get(content_type.lower(), "")
    if ext:
        return ext
    suffix = Path(name).suffix.lower()
    if suffix in {".png", ".jpg", ".jpeg", ".gif", ".bmp", ".tif", ".tiff", ".emf", ".wmf", ".svg", ".webp"}:
        return ".jpg" if suffix == ".jpeg" else ".tif" if suffix == ".tiff" else suffix
    return ".bin"


def _cell_text(cell) -> str:
    parts = [paragraph_text(p) for p in cell.paragraphs]
    return " / ".join(p for p in parts if p)


def _table_rows(table: Table) -> list[list[str]]:
    return [[_cell_text(cell) for cell in row.cells] for row in table.rows]


def _rows_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    norm = [row + [""] * (width - len(row)) for row in rows]

    def esc(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", " ")

    lines = [
        "| " + " | ".join(esc(cell) for cell in norm[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in norm[1:]:
        lines.append("| " + " | ".join(esc(cell) for cell in row) + " |")
    return "\n".join(lines)


def _iter_blocks(doc: Document):
    body = doc.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _paragraph_rids(block) -> list[str]:
    if isinstance(block, Paragraph):
        return _blip_ids(block._p)
    return []


def _caption_for(texts: list[str], blocks: list, index: int) -> tuple[str, int | None]:
    """Подпись принадлежит этому кадру, а не предыдущему рисунку."""
    nxt = index + 1
    if nxt < len(texts) and is_figure_label(texts[nxt]) and not _paragraph_rids(blocks[nxt]):
        return texts[nxt], nxt
    prev = index - 1
    if prev >= 0 and is_figure_label(texts[prev]) and not _paragraph_rids(blocks[prev]):
        earlier = prev - 1
        if earlier < 0 or not _paragraph_rids(blocks[earlier]):
            return texts[prev], prev
    if 0 <= index < len(texts) and is_figure_label(texts[index]):
        return texts[index], None
    return "", None


def _has_reference(texts: list[str], blocks: list, index: int) -> bool:
    for pos in (index + 1, index - 1, index):
        if not (0 <= pos < len(texts)):
            continue
        if pos != index and _paragraph_rids(blocks[pos]):
            continue
        if is_text_reference(texts[pos]) and not is_figure_label(texts[pos]):
            return True
    return False


@dataclass
class _Picked:
    rid: str
    caption: str
    reason: str
    block_index: int


def _pick_paragraph_images(texts: list[str], blocks: list, images: dict) -> tuple[list[_Picked], set[int]]:
    picked: list[_Picked] = []
    consumed: set[int] = set()
    for index, block in enumerate(blocks):
        if not isinstance(block, Paragraph):
            continue
        rids = [rid for rid in _blip_ids(block._p) if rid in images]
        if not rids:
            continue
        caption, caption_at = _caption_for(texts, blocks, index)
        if caption:
            reason = "caption"
            if caption_at is not None:
                consumed.add(caption_at)
        elif _has_reference(texts, blocks, index):
            reason = "reference"
            caption = ""
        else:
            continue
        for rid in rids:
            picked.append(_Picked(rid, caption, reason, index))
    return picked, consumed


def _pick_table_images(texts: list[str], blocks: list, images: dict) -> list[_Picked]:
    picked: list[_Picked] = []
    for index, block in enumerate(blocks):
        if not isinstance(block, Table):
            continue
        cells: list[tuple[str, list[str]]] = []
        for cell in block._tbl.iter(qn("w:tc")):
            rids = [rid for rid in _blip_ids(cell) if rid in images]
            if not rids:
                continue
            text = _norm(" ".join(node.text or "" for node in cell.findall(".//" + qn("w:t"))))
            cells.append((text, rids))
        if not cells:
            continue
        distinct = {rid for _, rids in cells for rid in rids}
        outside, _outside_at = _caption_for(texts, blocks, index)
        outside_ref = _has_reference(texts, blocks, index)
        for text, rids in cells:
            caption = text if is_figure_label(text) else ""
            if caption:
                reason = "caption"
            elif is_text_reference(text):
                reason = "reference"
            elif len(distinct) == 1 and outside:
                caption = outside
                reason = "caption"
            elif len(distinct) == 1 and outside_ref:
                reason = "reference"
            else:
                continue
            for rid in rids:
                picked.append(_Picked(rid, caption, reason, index))
    return picked


def convert_document(
    source: Path,
    *,
    extract_images: bool,
    ocr: Callable[[bytes, str], str] | None = None,
) -> ConvertResult:
    doc = load_document(source)
    dirname = media_dirname(source.stem)
    images = _load_images(doc) if extract_images else {}
    blocks = list(_iter_blocks(doc))
    texts = [paragraph_text(block) if isinstance(block, Paragraph) else "" for block in blocks]

    picked: list[_Picked] = []
    consumed: set[int] = set()
    if extract_images:
        para_picked, consumed = _pick_paragraph_images(texts, blocks, images)
        table_picked = _pick_table_images(texts, blocks, images)
        picked = para_picked + table_picked
        for item in table_picked:
            if item.caption:
                for pos, text in enumerate(texts):
                    if text == item.caption and abs(pos - item.block_index) == 1:
                        consumed.add(pos)

    seen_rids: dict[str, ImageRecord] = {}
    blobs: dict[str, bytes] = {}
    counter = 0
    picked_keys = {(item.block_index, item.rid) for item in picked}
    skipped = 0
    for index, block in enumerate(blocks):
        if isinstance(block, Paragraph):
            occurrence_ids = [rid for rid in _blip_ids(block._p) if rid in images]
        elif isinstance(block, Table):
            occurrence_ids = [rid for rid in _blip_ids(block._tbl) if rid in images]
        else:
            occurrence_ids = []
        skipped += sum(1 for rid in occurrence_ids if (index, rid) not in picked_keys)

    heading = ""
    md_lines: list[str] = []
    heading_count = 0
    table_count = 0

    def ensure_record(item: _Picked, current_heading: str) -> ImageRecord | None:
        nonlocal counter
        if item.rid in seen_rids:
            return seen_rids[item.rid]
        if item.rid not in images:
            return None
        name, content_type, blob = images[item.rid]
        counter += 1
        ext = _extension(name, content_type)
        file_name = f"img_{counter:03d}{ext}"
        record = ImageRecord(
            id=f"img_{counter:03d}",
            file=file_name,
            path=f"{dirname}/{file_name}",
            sha256=sha256_bytes(blob),
            nbytes=len(blob),
            heading=current_heading,
            caption=item.caption,
            reason=item.reason,
            preview=ext not in NO_PREVIEW,
        )
        if ocr is not None and record.preview:
            record_ocr = ocr(blob, ext)
            setattr(record, "ocr_text", record_ocr)
        seen_rids[item.rid] = record
        blobs[file_name] = blob
        return record

    picks_at: dict[int, list[_Picked]] = {}
    for item in picked:
        picks_at.setdefault(item.block_index, []).append(item)

    for index, block in enumerate(blocks):
        if index in consumed:
            continue
        if isinstance(block, Paragraph):
            for item in picks_at.get(index, []):
                record = ensure_record(item, heading)
                if record is not None:
                    _append_image(md_lines, record, item.caption)
            text = texts[index]
            level = heading_level(block) or (infer_heading_level(text) if text else None)
            if text and level:
                heading_count += 1
                heading = text
                md_lines.append(f"{'#' * level} {text}")
                md_lines.append("")
            elif text and not is_figure_label(text):
                prefix = "- " if _is_list(block) else ""
                md_lines.append(prefix + text)
                md_lines.append("")
            elif text and is_figure_label(text) and index not in {p.block_index for p in picked}:
                md_lines.append(text)
                md_lines.append("")
        elif isinstance(block, Table):
            table_count += 1
            rendered = _rows_to_markdown(_table_rows(block))
            if rendered:
                md_lines.append(rendered)
                md_lines.append("")
            for item in picks_at.get(index, []):
                record = ensure_record(item, heading)
                if record is not None:
                    _append_image(md_lines, record, item.caption)

    records = list(seen_rids.values())
    return ConvertResult(
        markdown="\n".join(md_lines).rstrip() + "\n",
        media_dirname=dirname,
        images=records,
        blobs=blobs,
        saved=len(records),
        skipped=skipped,
        headings=heading_count,
        tables=table_count,
    )


def _append_image(lines: list[str], record: ImageRecord, caption: str) -> None:
    alt = caption or record.file
    lines.append(f"![{_md_alt(alt)}]({record.path})")
    lines.append("")
    if caption:
        lines.append(f"*{caption}*")
        lines.append("")


def _md_alt(text: str) -> str:
    return text.replace("[", "\\[").replace("]", "\\]")


def image_ocr_text(record: ImageRecord) -> str | None:
    if not hasattr(record, "ocr_text"):
        return None
    return getattr(record, "ocr_text")
