"""XLSX → оглавление Markdown, карта листов и CSV по каждому видимому листу."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

_INVALID_NAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')
PREVIEW_ROWS = 30


@dataclass
class FormulaNote:
    address: str
    formula: str
    cached: bool


@dataclass
class SheetGrid:
    name: str
    hidden: bool
    rows: list[list[str]] = field(default_factory=list)
    header_row: int | None = None
    merges: list[str] = field(default_factory=list)
    formulas: list[FormulaNote] = field(default_factory=list)
    csv_name: str | None = None


@dataclass
class WorkbookData:
    source_name: str
    sheets: list[SheetGrid]


def sheet_filename(name: str, used: set[str]) -> str:
    cleaned = _INVALID_NAME.sub("_", name).strip(" .") or "sheet"
    candidate = cleaned
    number = 2
    while candidate.lower() in used:
        candidate = f"{cleaned}_{number}"
        number += 1
    used.add(candidate.lower())
    return f"{candidate}.csv"


def format_cell_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0, 0):
            return value.date().isoformat()
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value)


def read_workbook(
    source: Path,
    *,
    visible_only: bool = True,
    repeat_merges: bool = True,
    keep_formulas: bool = True,
) -> WorkbookData:
    from openpyxl import load_workbook

    formula_book = load_workbook(source, data_only=False)
    cached_book = load_workbook(source, data_only=True)
    try:
        sheets = [
            _read_sheet(
                formula_book[name],
                cached_book[name],
                visible_only=visible_only,
                repeat_merges=repeat_merges,
                keep_formulas=keep_formulas,
            )
            for name in formula_book.sheetnames
        ]
    finally:
        formula_book.close()
        cached_book.close()
    _assign_csv_names(sheets, visible_only=visible_only)
    return WorkbookData(source_name=source.name, sheets=sheets)


def render_workbook(data: WorkbookData, folder_name: str) -> tuple[str, dict[str, str], dict]:
    """Возвращает Markdown, содержимое CSV и тело sheets.json."""
    lines = [f"# {data.source_name}", ""]
    csv_files: dict[str, str] = {}
    meta_sheets: list[dict] = []
    for sheet in data.sheets:
        lines.append(f"## {sheet.name}")
        lines.append("")
        if sheet.hidden and sheet.csv_name is None:
            lines.append("Лист скрыт и пропущен.")
            lines.append("")
            meta_sheets.append({"name": sheet.name, "hidden": True})
            continue
        row_count = len(sheet.rows)
        col_count = max((len(row) for row in sheet.rows), default=0)
        header = "не определена" if sheet.header_row is None else str(sheet.header_row)
        lines.append(f"Строк: {row_count}, столбцов: {col_count}. Строка заголовков: {header}.")
        lines.append("")
        rel = f"{folder_name}/{sheet.csv_name}"
        lines.append(f"Полные данные: [{sheet.csv_name}]({rel})")
        lines.append("")
        if row_count > PREVIEW_ROWS:
            lines.append(f"Показаны первые {PREVIEW_ROWS} строк. Полные данные в CSV.")
            lines.append("")
        preview = _rows_to_markdown(sheet.rows[:PREVIEW_ROWS])
        if preview:
            lines.append(preview)
            lines.append("")
        elif row_count == 0:
            lines.append("Лист пуст.")
            lines.append("")
        if sheet.csv_name:
            csv_files[sheet.csv_name] = _rows_to_csv(sheet.rows)
        entry: dict[str, Any] = {
            "name": sheet.name,
            "hidden": sheet.hidden,
            "rows": row_count,
            "cols": col_count,
            "header_row": sheet.header_row,
            "csv": rel,
            "merges": sheet.merges,
            "formulas": [
                {"address": item.address, "formula": item.formula, "cached": item.cached}
                for item in sheet.formulas
            ],
        }
        meta_sheets.append(entry)
    payload = {"source": data.source_name, "sheets": meta_sheets}
    return "\n".join(lines).rstrip() + "\n", csv_files, payload


def iter_text_slots(data: WorkbookData):
    """Пары (куда записать, текст) для деперсонализации ячеек и формул."""
    for sheet in data.sheets:
        if sheet.csv_name is None:
            continue
        for row in sheet.rows:
            for index, value in enumerate(row):
                if value:
                    yield ("cell", sheet, row, index), value
        for note in sheet.formulas:
            if note.formula:
                yield ("formula", note), note.formula


def write_slot(slot, value: str) -> None:
    kind = slot[0]
    if kind == "cell":
        _kind, _sheet, row, index = slot
        row[index] = value
    else:
        _kind, note = slot
        note.formula = value


def _read_sheet(formula_ws, cached_ws, *, visible_only: bool, repeat_merges: bool, keep_formulas: bool) -> SheetGrid:
    hidden = formula_ws.sheet_state != "visible"
    sheet = SheetGrid(name=formula_ws.title, hidden=hidden)
    if hidden and visible_only:
        return sheet
    merges = list(formula_ws.merged_cells.ranges)
    sheet.merges = [str(item) for item in merges]
    bounds = _bounds(formula_ws, merges)
    if bounds is None:
        return sheet
    min_row, min_col, max_row, max_col = bounds
    grid: list[list[str]] = []
    for row_idx in range(min_row, max_row + 1):
        line: list[str] = []
        for col_idx in range(min_col, max_col + 1):
            formula_cell = formula_ws.cell(row_idx, col_idx)
            cached_cell = cached_ws.cell(row_idx, col_idx)
            text, note = _cell_export(formula_cell, cached_cell, keep_formulas=keep_formulas)
            line.append(text)
            if note is not None:
                sheet.formulas.append(note)
        grid.append(line)
    if repeat_merges:
        _repeat_merges(grid, merges, min_row, min_col)
    sheet.rows = grid
    sheet.header_row = _header_row(grid, merges, min_row, min_col)
    return sheet


def _cell_export(formula_cell, cached_cell, *, keep_formulas: bool) -> tuple[str, FormulaNote | None]:
    if formula_cell.data_type == "f":
        formula = str(formula_cell.value or "")
        cached_value = cached_cell.value
        cached = cached_value is not None or cached_cell.data_type == "e"
        text = "" if not cached else format_cell_value(cached_value if cached_cell.data_type != "e" else cached_cell.value)
        if cached_cell.data_type == "e":
            text = format_cell_value(cached_cell.value)
            cached = True
        note = None
        if keep_formulas:
            note = FormulaNote(
                address=formula_cell.coordinate,
                formula=formula if formula.startswith("=") else f"={formula}",
                cached=cached,
            )
        if not cached:
            text = ""
        return text, note
    if formula_cell.data_type == "e":
        return format_cell_value(formula_cell.value), None
    return format_cell_value(formula_cell.value), None


def _bounds(ws, merges) -> tuple[int, int, int, int] | None:
    rows: list[int] = []
    cols: list[int] = []
    for (row_idx, col_idx), cell in ws._cells.items():
        if cell.value is None and cell.data_type != "e":
            continue
        rows.append(row_idx)
        cols.append(col_idx)
    for merged in merges:
        rows.extend((merged.min_row, merged.max_row))
        cols.extend((merged.min_col, merged.max_col))
    if not rows:
        return None
    return min(rows), min(cols), max(rows), max(cols)


def _repeat_merges(grid: list[list[str]], merges, min_row: int, min_col: int) -> None:
    for merged in merges:
        top = merged.min_row - min_row
        left = merged.min_col - min_col
        if not (0 <= top < len(grid) and 0 <= left < len(grid[top])):
            continue
        value = grid[top][left]
        for row_idx in range(merged.min_row, merged.max_row + 1):
            for col_idx in range(merged.min_col, merged.max_col + 1):
                rr = row_idx - min_row
                cc = col_idx - min_col
                if 0 <= rr < len(grid) and 0 <= cc < len(grid[rr]):
                    grid[rr][cc] = value


def _header_row(rows: list[list[str]], merges, min_row: int, min_col: int) -> int | None:
    width = max((len(row) for row in rows), default=0)
    full_width_rows = _full_width_merge_rows(merges, width, min_col)
    for offset, row in enumerate(rows):
        filled = sum(1 for cell in row if cell.strip())
        excel_row = min_row + offset
        if excel_row in full_width_rows:
            continue
        if filled < 2:
            continue
        return excel_row
    return None


def _full_width_merge_rows(merges, width: int, min_col: int) -> set[int]:
    found: set[int] = set()
    if width < 2:
        return found
    last_col = min_col + width - 1
    for merged in merges:
        if merged.min_row != merged.max_row:
            continue
        if merged.min_col <= min_col and merged.max_col >= last_col:
            found.add(merged.min_row)
    return found


def _assign_csv_names(sheets: list[SheetGrid], *, visible_only: bool) -> None:
    used: set[str] = set()
    for sheet in sheets:
        if sheet.hidden and visible_only:
            continue
        sheet.csv_name = sheet_filename(sheet.name, used)


def _rows_to_csv(rows: list[list[str]]) -> str:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerows(rows)
    return buffer.getvalue()


def _rows_to_markdown(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    norm = [row + [""] * (width - len(row)) for row in rows]

    def esc(value: str) -> str:
        return value.replace("|", "\\|").replace("\n", " ")

    lines = [
        "| " + " | ".join(esc(cell) for cell in norm[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    for row in norm[1:]:
        lines.append("| " + " | ".join(esc(cell) for cell in row) + " |")
    return "\n".join(lines)

