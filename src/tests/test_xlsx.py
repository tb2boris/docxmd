import csv
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

from docxpipe.pipeline import convert_xlsx_tree
from docxpipe.xlsx_convert import format_cell_value


def _book(path: Path) -> Path:
    wb = Workbook()
    first = wb.active
    first.title = "Реестр"
    first.merge_cells("A1:B1")
    first["A1"] = "Общий заголовок"
    first["A2"] = "Имя"
    first["B2"] = "Сумма"
    first["A3"] = "Иван"
    first["B3"] = 12.0
    first["A4"] = None
    first["B4"] = None
    first["A5"] = datetime(2024, 3, 1, 13, 45, 0)
    first["B5"] = "=B3+1"
    second = wb.create_sheet("A<B")
    second["A1"] = "лево"
    second["B1"] = "право"
    twin = wb.create_sheet("A_B")
    twin["A1"] = "один"
    twin["B1"] = "два"
    hidden = wb.create_sheet("Секрет")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "нельзя"
    hidden["B1"] = "видеть"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def test_format_cell_value_iso_and_errors():
    assert format_cell_value(datetime(2024, 3, 1)) == "2024-03-01"
    assert format_cell_value(datetime(2024, 3, 1, 13, 45, 0)) == "2024-03-01T13:45:00"
    assert format_cell_value(12.0) == "12"
    assert format_cell_value("#N/A") == "#N/A"
    assert format_cell_value("#REF!") == "#REF!"


def test_workbook_outline_csv_and_hidden(tmp_path: Path):
    source = _book(tmp_path / "реестр.xlsx")
    before = source.read_bytes()
    out = tmp_path / "out"
    outcomes = convert_xlsx_tree(source, out, overwrite=True)
    assert outcomes[0].error is None
    text = (out / "реестр.md").read_text(encoding="utf-8")
    payload = json.loads((out / "реестр.sheets.json").read_text(encoding="utf-8"))
    assert "Лист скрыт и пропущен." in text
    assert "Секрет" in text
    assert "[Реестр.csv](реестр/Реестр.csv)" in text
    assert "A_B_2.csv" in text or "A_B.csv" in text
    hidden = next(item for item in payload["sheets"] if item["name"] == "Секрет")
    assert hidden == {"name": "Секрет", "hidden": True}
    visible = next(item for item in payload["sheets"] if item["name"] == "Реестр")
    assert "A1:B1" in visible["merges"]
    assert visible["header_row"] == 2
    formula = next(item for item in visible["formulas"] if item["address"] == "B5")
    assert formula["cached"] is False
    assert formula["formula"] == "=B3+1"
    raw = (out / "реестр" / "Реестр.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(raw.decode("utf-8-sig").splitlines()))
    assert rows[0] == ["Общий заголовок", "Общий заголовок"]
    assert rows[1] == ["Имя", "Сумма"]
    assert rows[2] == ["Иван", "12"]
    assert rows[3] == ["", ""]
    assert rows[4][0] == "2024-03-01T13:45:00"
    assert rows[4][1] == ""
    names = sorted(item.name for item in (out / "реестр").iterdir())
    assert "Секрет.csv" not in names
    assert "A_B.csv" in names
    assert "A_B_2.csv" in names
    assert source.read_bytes() == before


def test_duplicate_workbook_stem(tmp_path: Path):
    left = tmp_path / "a"
    right = tmp_path / "b"
    _book(left / "реестр.xlsx")
    _book(right / "реестр.xlsx")
    out = tmp_path / "out"
    outcomes = convert_xlsx_tree(tmp_path, out, overwrite=True)
    assert all(item.error is None for item in outcomes)
    assert (out / "реестр.md").is_file()
    assert (out / "реестр_2.md").is_file()
    assert (out / "реестр").is_dir()
    assert (out / "реестр_2").is_dir()


def test_depersonalize_cells_keeps_sheet_name_and_link(tmp_path: Path):
    wb = Workbook()
    sheet = wb.active
    sheet.title = "Контакты"
    sheet["A1"] = "Почта"
    sheet["B1"] = "Комментарий"
    sheet["A2"] = "user@example.com"
    sheet["B2"] = "рабочая"
    source = tmp_path / "book.xlsx"
    wb.save(source)
    out = tmp_path / "out"
    outcomes = convert_xlsx_tree(
        source,
        out,
        depersonalize=True,
        save_report=True,
        overwrite=True,
        seed=1,
    )
    assert outcomes[0].error is None
    csv_path = out / "book" / "Контакты.csv"
    body = csv_path.read_text(encoding="utf-8-sig")
    markdown = (out / "book.md").read_text(encoding="utf-8")
    report = (out / "book.report.json").read_text(encoding="utf-8")
    assert "user@example.com" not in body
    assert "user@example.com" not in markdown
    assert "user@example.com" not in report
    assert "[Контакты.csv](book/Контакты.csv)" in markdown
    assert csv_path.is_file()
