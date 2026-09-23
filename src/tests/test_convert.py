import base64
from pathlib import Path

from docx import Document
from docx.shared import Inches

import json

from docxpipe.anonymize import anonymize_markdown, apply_restore, begin_pair_log, mask_image_links, restore_image_links
from docxpipe.convert import convert_document, is_figure_label, is_text_reference, media_dirname
from docxpipe.pipeline import convert_tree, restore_file

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def _png(path: Path) -> Path:
    path.write_bytes(PNG)
    return path


def _save(doc: Document, path: Path) -> Path:
    doc.save(path)
    return path


def test_label_and_reference_patterns():
    assert is_figure_label("Рисунок 1 — Контур")
    assert is_figure_label("Схема 2. Поток")
    assert is_figure_label("Рис. 3")
    assert not is_figure_label("Логотип компании")
    assert is_text_reference("как показано на схеме")
    assert is_text_reference("см. рис. 4 в разделе")
    assert not is_text_reference("Логотип компании")


def test_media_dirname_is_bound_to_stem():
    assert media_dirname("Регламент") == "image_Регламент"
    assert media_dirname('a:b*c') == "image_a_b_c"


def test_structure_and_figure_selection(tmp_path: Path):
    png = _png(tmp_path / "pixel.png")
    doc = Document()
    doc.add_heading("Раздел 1", level=1)
    doc.add_paragraph("Обычный абзац.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Код"
    table.cell(0, 1).text = "Имя"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "Насос"
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("Рисунок 1 — Контур")
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("Логотип компании")
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("как показано на схеме")
    source = _save(doc, tmp_path / "регламент.docx")

    result = convert_document(source, extract_images=True)
    text = result.markdown
    assert "# Раздел 1" in text
    assert "| Код | Имя |" in text
    assert "| 1 | Насос |" in text
    assert "image_регламент/img_001.png" in text
    assert "Рисунок 1 — Контур" in text
    assert "Логотип компании" in text
    assert text.count("![") == 2
    assert "как показано на схеме" in text
    assert result.saved == 1
    assert result.skipped == 1
    assert result.images[0].reason == "caption"
    assert all(item.file.startswith("img_") for item in result.images)
    assert "ocr_text" not in result.images[0].as_dict()


def test_two_files_get_separate_image_folders(tmp_path: Path):
    png = _png(tmp_path / "pixel.png")
    src = tmp_path / "in"
    a = src / "a"
    b = src / "b"
    a.mkdir(parents=True)
    b.mkdir()
    for folder, caption in ((a, "Рисунок 1 — А"), (b, "Схема 1 — Б")):
        doc = Document()
        doc.add_picture(str(png), width=Inches(1))
        doc.add_paragraph(caption)
        doc.save(folder / "карточка.docx")

    out = tmp_path / "out"
    outcomes = convert_tree(
        src,
        out,
        extract_images=True,
        ocr=False,
        depersonalize=False,
        save_report=False,
        seed=None,
        overwrite=True,
    )
    assert all(item.error is None for item in outcomes)
    folders = sorted(path.name for path in out.iterdir() if path.is_dir())
    assert folders == ["image_карточка", "image_карточка_2"]
    assert (out / "image_карточка" / "images.json").is_file()
    assert (out / "image_карточка_2" / "images.json").is_file()
    md_names = sorted(path.name for path in out.glob("*.md"))
    assert md_names == ["карточка.md", "карточка_2.md"]


def test_ocr_flag_off_omits_ocr_text(tmp_path: Path):
    png = _png(tmp_path / "pixel.png")
    doc = Document()
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("Диаграмма 1")
    source = _save(doc, tmp_path / "diagram.docx")
    out = tmp_path / "out"
    outcomes = convert_tree(
        source,
        out,
        extract_images=True,
        ocr=False,
        depersonalize=False,
        save_report=False,
        seed=None,
        overwrite=True,
    )
    assert outcomes[0].error is None
    payload = (out / "image_diagram" / "images.json").read_text(encoding="utf-8")
    assert "ocr_text" not in payload


def test_image_links_survive_anonymizer():
    source = "Контакт user@example.com\n\n![Рисунок 1](image_doc/img_001.png)\n"
    masked, tokens = mask_image_links(source)
    assert "image_doc" not in masked
    assert tokens == ["![Рисунок 1](image_doc/img_001.png)"]

    class _Fake:
        def process_text(self, text: str, dry_run: bool = False):
            replaced = text.replace("image_", "LEAK_").replace("user@example.com", "anon@example.com")
            records = [{
                "category": "email",
                "original_hash": "sha256:abc",
                "replacement": "anon@example.com",
            }]
            return replaced, records

    updated, records = anonymize_markdown(_Fake(), source)
    assert "![Рисунок 1](image_doc/img_001.png)" in updated
    assert "anon@example.com" in updated
    assert "LEAK_" not in updated
    assert records[0]["category"] == "email"
    assert restore_image_links(masked, tokens) == source


def test_depersonalize_replaces_email_and_keeps_image_link(tmp_path: Path):
    png = _png(tmp_path / "pixel.png")
    doc = Document()
    doc.add_paragraph("Почта user@example.com")
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("Рисунок 1 — Схема")
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("Логотип")
    source = _save(doc, tmp_path / "card.docx")
    out = tmp_path / "out"
    outcomes = convert_tree(
        source,
        out,
        extract_images=True,
        ocr=False,
        depersonalize=True,
        save_report=True,
        seed=1,
        overwrite=True,
    )
    assert outcomes[0].error is None
    text = (out / "card.md").read_text(encoding="utf-8")
    report = (out / "card.report.json").read_text(encoding="utf-8")
    assert "user@example.com" not in text
    assert "user@example.com" not in report
    assert "image_card/img_001.png" in text
    assert (out / "image_card" / "img_001.png").is_file()
    assert not (out / "image_card" / "img_002.png").exists()
    assert outcomes[0].detections >= 1


def test_restore_map_roundtrip_keeps_image_link(tmp_path: Path):
    png = _png(tmp_path / "pixel.png")
    doc = Document()
    doc.add_paragraph("Почта user@example.com")
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("Рисунок 1 — Схема")
    source = _save(doc, tmp_path / "card.docx")
    out = tmp_path / "out"
    outcomes = convert_tree(
        source,
        out,
        extract_images=True,
        ocr=False,
        depersonalize=True,
        save_report=True,
        save_restore_map=True,
        seed=1,
        overwrite=True,
    )
    assert outcomes[0].error is None
    assert outcomes[0].restore_path is not None
    text = (out / "card.md").read_text(encoding="utf-8")
    report = (out / "card.report.json").read_text(encoding="utf-8")
    payload = json.loads((out / "card.restore.json").read_text(encoding="utf-8"))
    pair = next(item for item in payload["pairs"] if item["original"] == "user@example.com")
    assert pair["replacement"]
    assert pair["replacement"] in text
    assert "user@example.com" not in text
    assert "user@example.com" not in report
    assert "image_card/img_001.png" in text
    edited = text + f"\n\nПовтор {pair['replacement']}.\n"
    edited_path = tmp_path / "edited.md"
    edited_path.write_text(edited, encoding="utf-8")
    restored = restore_file(
        edited_path,
        out / "card.restore.json",
        tmp_path / "edited.restored.md",
        overwrite=True,
    )
    assert restored.error is None
    body = restored.output_path.read_text(encoding="utf-8")
    assert body.count("user@example.com") == 2
    assert pair["replacement"] not in body
    assert "image_card/img_001.png" in body


def test_restore_map_is_not_written_when_disabled(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Почта user@example.com")
    source = _save(doc, tmp_path / "card.docx")
    out = tmp_path / "out"
    outcomes = convert_tree(
        source,
        out,
        extract_images=False,
        ocr=False,
        depersonalize=True,
        save_report=False,
        save_restore_map=False,
        seed=1,
        overwrite=True,
    )
    assert outcomes[0].error is None
    assert outcomes[0].restore_path is None
    assert not list(out.glob("*.restore.json"))


def test_apply_restore_skips_inflected_and_ambiguous_forms():
    text, applied, missing, ambiguous = apply_restore(
        "Кузнецова пришла. Кузнецов Пётр Сергеевич тоже. ![x](image_card/img_001.png)",
        [{
            "category": "fio",
            "original": "Иванов Иван Иванович",
            "replacement": "Кузнецов Пётр Сергеевич",
        }],
    )
    kept, applied_short, missing_short, _ambiguous = apply_restore(
        "Кузнецова пришла",
        [{"category": "fio", "original": "Иванов", "replacement": "Кузнецов"}],
    )
    assert kept == "Кузнецова пришла"
    assert applied_short == []
    assert len(missing_short) == 1

    assert "Кузнецова пришла." in text
    assert "Иванов Иван Иванович тоже." in text
    assert "image_card/img_001.png" in text
    assert len(applied) == 1
    assert missing == []
    assert ambiguous == []

    untouched, applied, missing, ambiguous = apply_restore(
        "Кузнецов здесь",
        [
            {"category": "fio", "original": "Иванов", "replacement": "Кузнецов"},
            {"category": "fio", "original": "Петров", "replacement": "Кузнецов"},
        ],
    )
    assert untouched == "Кузнецов здесь"
    assert applied == []
    assert len(ambiguous) == 2


def test_pair_log_records_original_during_replacement():
    class _Mapper:
        def get_replacement(self, original: str, category: str, generator_name: str | None = None) -> str:
            return "Кузнецов Пётр Сергеевич"

    class _Processor:
        def __init__(self) -> None:
            self.mapper = _Mapper()

        def process_text(self, text: str, dry_run: bool = False):
            replacement = self.mapper.get_replacement("Иванов Иван Иванович", "fio", None)
            return text.replace("Иванов Иван Иванович", replacement), []

    processor = _Processor()
    log = begin_pair_log(processor)
    updated, _records = anonymize_markdown(processor, "Иванов Иван Иванович")
    assert updated == "Кузнецов Пётр Сергеевич"
    assert log.pairs == [{
        "category": "fio",
        "original": "Иванов Иван Иванович",
        "replacement": "Кузнецов Пётр Сергеевич",
    }]


def test_ocr_text_is_stored_when_callback_returns_text(tmp_path: Path):
    png = _png(tmp_path / "pixel.png")
    doc = Document()
    doc.add_picture(str(png), width=Inches(1))
    doc.add_paragraph("График 1")
    source = _save(doc, tmp_path / "graph.docx")
    result = convert_document(source, extract_images=True, ocr=lambda blob, ext: "Иванов Иван")
    assert result.images[0].as_dict(ocr_text="Иванов Иван")["ocr_text"] == "Иванов Иван"
